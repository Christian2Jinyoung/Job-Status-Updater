import os
import re
from typing import Literal

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.worksheet.worksheet import Worksheet

from config import TRACKER_FILE, SHEET_NAME, COLUMNS, UNMATCHED_FILE, UNMATCHED_COLUMNS


def _load_or_create(path: str, columns: list[str]) -> openpyxl.Workbook:
    if os.path.exists(path):
        return openpyxl.load_workbook(path)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(columns)
    return workbook


COMPANY_SUFFIXES = {"inc", "incorporated", "corp", "corporation", "co", "company", "llc", "ltd", "limited", "plc"}


def _company_tokens(name: str) -> list[str]:
    tokens = re.sub(r"[^a-z0-9]+", " ", name.lower()).split()
    while tokens and tokens[-1] in COMPANY_SUFFIXES:
        tokens.pop()
    return tokens


def _companies_match(a: list[str], b: list[str]) -> bool:
    """True if one name is the other plus extra trailing words (e.g. Allstate / Allstate Insurance)."""
    shorter, longer = sorted((a, b), key=len)
    return bool(shorter) and longer[: len(shorter)] == shorter


def _find_row(sheet: Worksheet, company: str, position: str):
    company_idx = COLUMNS.index("Company")
    position_idx = COLUMNS.index("Position")
    target = _company_tokens(company)

    exact, partial = [], []
    for row in sheet.iter_rows(min_row=2):
        row_company = row[company_idx].value
        row_position = row[position_idx].value
        if not (row_company and row_position):
            continue
        if str(row_position).strip().lower() != position.strip().lower():
            continue
        row_tokens = _company_tokens(str(row_company))
        if row_tokens == target:
            exact.append(row)
        elif _companies_match(row_tokens, target):
            partial.append(row)

    if exact:
        return exact[0]
    # Only trust a loose match when it's unambiguous.
    if len(partial) == 1:
        return partial[0]
    return None


def update_entry(
    status: str | None,
    company: str,
    position: str,
    location: str,
    job_type: str,
    email_timestamp: str,
    notes: str,
) -> Literal["updated", "duplicate", "unmatched"]:
    """Update the existing tracker row for this Company+Position.

    Date Applied is never touched; the email's received timestamp is prefixed to the note instead.
    A row whose notes already have a line starting with this timestamp is treated
    as already processed and left alone.
    """
    workbook = _load_or_create(TRACKER_FILE, COLUMNS)
    sheet = workbook[SHEET_NAME]

    existing_row = _find_row(sheet, company, position)
    if not existing_row:
        return "unmatched"

    col = {name: i for i, name in enumerate(COLUMNS)}
    notes_cell = existing_row[col["Notes"]]
    existing_notes = str(notes_cell.value or "")
    if any(line.startswith(email_timestamp) for line in existing_notes.splitlines()):
        return "duplicate"

    if status:
        existing_row[col["Status"]].value = status
    # "Unknown" would clobber a blank or a value the user filled in by hand.
    if location != "Unknown":
        existing_row[col["Location"]].value = location
    if job_type != "Unknown":
        existing_row[col["Type"]].value = job_type
    new_note = f"{email_timestamp} {notes}"
    notes_cell.value = f"{new_note}\n{existing_notes}" if existing_notes else new_note
    notes_cell.alignment = Alignment(wrap_text=True, vertical="top")

    workbook.save(TRACKER_FILE)
    return "updated"


def append_unmatched(update: dict, subject: str, sender: str, received_date: str) -> bool:
    """Record an unmatched email in the unmatched file. Returns False if it was already recorded."""
    workbook = _load_or_create(UNMATCHED_FILE, UNMATCHED_COLUMNS)
    sheet = workbook[SHEET_NAME]

    sender_idx = UNMATCHED_COLUMNS.index("Sender")
    received_idx = UNMATCHED_COLUMNS.index("Received")
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[sender_idx] == sender and row[received_idx] == received_date:
            return False

    sheet.append([
        update["Status"],
        update["Company"],
        update["Position"],
        None if update["Location"] == "Unknown" else update["Location"],
        None if update["Type"] == "Unknown" else update["Type"],
        update["Date"],
        update["Notes"],
        subject,
        sender,
        received_date,
    ])
    workbook.save(UNMATCHED_FILE)
    return True
